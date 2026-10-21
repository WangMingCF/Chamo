from pathlib import Path
from reportlab.pdfgen import canvas
from PIL import Image,PngImagePlugin
from pptx import Presentation
from pptx.util import Inches
root=Path(__file__).resolve().parents[1]/'tests/fixtures';root.mkdir(exist_ok=True)
c=canvas.Canvas(str(root/'示例讲义.pdf'))
for n in range(6):
    c.setFont('Helvetica',28);c.drawString(55,760,f'Chamo / Page {n+1}')
    c.setFont('Helvetica',18);c.drawString(55,700,'2 + 3 = 5');c.rect(55,400,300,220,fill=0);c.showPage()
c.save()
p=Presentation()
for title,text in [('课程示例 / Lesson','数学：2 + 3 = 5'),('Table and image','Native Office rendering')]:
    s=p.slides.add_slide(p.slide_layouts[5]);s.shapes.title.text=title;s.shapes.add_textbox(Inches(1),Inches(2),Inches(7), Inches(1)).text=text
    s.notes_slide.notes_text_frame.text='Notes deliberately remain in the original PPTX.'
p.save(root/'示例课件.pptx')
im=Image.new('RGBA',(800,600),(40,100,200,150));info=PngImagePlugin.PngInfo();info.add_text('Description','Synthetic RGBA sample');im.save(root/'透明图片.png',pnginfo=info,compress_level=0)
(root/'损坏样例.pdf').write_bytes(b'Intentionally invalid synthetic PDF')
