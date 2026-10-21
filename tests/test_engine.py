import sys,json,subprocess,hashlib,struct,zlib,threading,time
from pathlib import Path
import pytest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'engine'))
import workbench as w
import platform_office
from PIL import Image,PngImagePlugin
from reportlab.pdfgen import canvas
ROOT=Path(__file__).resolve().parents[1]
@pytest.fixture
def cfg():
    # Resolved per platform so the same suite runs on macOS and Linux; Office
    # tests skip when no renderer is installed instead of failing on a path.
    return {'operation':'pdf_images','dpi':72,'format':'png','quality':90,'range':'','pdf_pages':100,'pdf_bytes':100000000,'workers':2,'max_edge':0,'batch_files':50,'lossy':False,'soffice':platform_office.find(ROOT/'engine') or ''}

def office_available():
    return platform_office.available(ROOT/'engine')
def pdf(path,n=3):
    c=canvas.Canvas(str(path))
    for i in range(n):c.drawString(80,700,f'PAGE {i+1} / source');c.showPage()
    c.save();return path

def run(source,out,cfg):return w.one({'sources':[str(source)],'output':str(out),'relative':source.name,'config':cfg})

def test_pages():
    assert w.pages('1-3,5',5)==[0,1,2,4]
    for s in ['0','1-99','1,1','x','2-1','1-2-3']:
        with pytest.raises(ValueError):w.pages(s,5)

def test_render_hash_resume_and_corruption(tmp_path,cfg):
    src=pdf(tmp_path/'中文 数学.pdf');before=w.sha(src);cfg['range']='1,3'
    r=run(src,tmp_path/'out',cfg)
    assert len(r['outputs'])==2 and w.sha(src)==before
    assert run(src,tmp_path/'out',cfg)['state']=='resumed'
    Path(r['outputs'][0]['path']).write_bytes(b'changed')
    with pytest.raises(ValueError,match='改变'):run(src,tmp_path/'out',cfg)

def test_strict_split_and_rag_native(tmp_path,cfg):
    src=pdf(tmp_path/'split.pdf',7);cfg['operation']='pdf_split';cfg['pdf_pages']=3
    r=run(src,tmp_path/'out',cfg)
    assert sorted(x['pages'] for x in r['outputs'])==[1,3,3]
    assert all(x['bytes']<=cfg['pdf_bytes'] for x in r['outputs'])
    assert sum(len(w.reader(x['path']).pages) for x in r['outputs'])==7
    from pptx import Presentation
    ppt=tmp_path/'保留.pptx';p=Presentation();p.slides.add_slide(p.slide_layouts[0]);p.save(ppt)
    cfg['operation']='rag';r=run(ppt,tmp_path/'rag',cfg)
    assert w.sha(r['outputs'][0]['path'])==w.sha(ppt)

def test_png_ancillary_lossless(tmp_path,cfg):
    src=tmp_path/'meta.png';im=Image.new('RGBA',(120,80),(40,90,100,100));meta=PngImagePlugin.PngInfo();meta.add_text('meaning','retain me');im.save(src,pnginfo=meta,compress_level=0)
    cfg['operation']='image_compress';r=run(src,tmp_path/'out',cfg)
    with Image.open(r['outputs'][0]['path']) as b:assert b.info['meaning']=='retain me';assert b.tobytes()==im.tobytes()
    assert r['outputs'][0]['bytes']<src.stat().st_size

def test_multiframe_rejected(tmp_path,cfg):
    src=tmp_path/'animated.gif';a=Image.new('RGB',(20,20),'red');b=Image.new('RGB',(20,20),'blue');a.save(src,save_all=True,append_images=[b])
    cfg['operation']='image_convert'
    with pytest.raises(ValueError,match='多帧'):run(src,tmp_path/'out',cfg)

def test_encrypted_and_corrupt(tmp_path,cfg):
    from pypdf import PdfWriter
    src=tmp_path/'encrypted.pdf';p=PdfWriter();p.add_blank_page(100,100);p.encrypt('secret');p.write(src)
    with pytest.raises(ValueError,match='加密'):run(src,tmp_path/'out',cfg)
    src=tmp_path/'broken.pdf';src.write_bytes(b'not a pdf')
    with pytest.raises(Exception):run(src,tmp_path/'out',cfg)

def test_merge_and_images(tmp_path,cfg):
    a=pdf(tmp_path/'a.pdf',2);b=pdf(tmp_path/'b.pdf',3);cfg['operation']='pdf_merge'
    r=w.one({'sources':[str(a),str(b)],'output':str(tmp_path/'out'),'config':cfg})
    assert r['outputs'][0]['pages']==5
    img=tmp_path/'alpha.png';Image.new('RGBA',(50,100),(20,30,40,120)).save(img)
    cfg['operation']='images_pdf';r=run(img,tmp_path/'imgout',cfg);assert r['outputs'][0]['pages']==1

def test_scan_scope(tmp_path):
    inp=tmp_path/'资料';inp.mkdir();(inp/'x.ts').write_bytes(b'x');(inp/'a.txt').write_text('a')
    # Creating a symlink needs elevation or developer mode on Windows. The scan
    # still has to skip one, so only create it where that is possible.
    try:
        (inp/'link.txt').symlink_to(inp/'a.txt')
    except (OSError, NotImplementedError):
        pass
    out=inp/'output';out.mkdir();(out/'old.pdf').write_bytes(b'x')
    rows,n=w.scan([str(inp)],out.resolve(),'rag');assert n==1;assert len(rows)==1;assert rows[0]['relative']=='资料/a.txt'
    with pytest.raises(ValueError):w.scan([str(inp)],inp.resolve(),'rag')

@pytest.mark.skipif(not office_available(), reason="no LibreOffice renderer on this platform")
def test_office_render(tmp_path,cfg):
    from pptx import Presentation
    from pptx.util import Inches
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[5]);s.shapes.title.text='Math lesson';s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text='2 + 3 = 5'
    src=tmp_path/'课程.pptx';p.save(src);cfg['operation']='office_pdf';r=run(src,tmp_path/'out',cfg)
    assert r['outputs'][0]['pages']==1
    assert '2 + 3' in w.reader(r['outputs'][0]['path']).pages[0].extract_text()

def test_supervisor_failure_isolation(tmp_path,cfg):
    a=pdf(tmp_path/'good.pdf');b=tmp_path/'bad.pdf';b.write_bytes(b'bad')
    req=tmp_path/'request.json';req.write_text(json.dumps({'inputs':[str(a),str(b)],'output':str(tmp_path/'out'),'config':cfg}))
    r=subprocess.run([sys.executable,'-I',str(ROOT/'engine/workbench.py'),'--request',str(req)],input='',capture_output=True,text=True,encoding='utf-8',timeout=180)
    events=[json.loads(x) for x in r.stdout.splitlines()];final=events[-1]
    assert final['type']=='finished';assert final['success']==1;assert final['failed']==1
    assert len(list((Path(final['report'])).glob('batch-*.json')))==1

def test_scan_skips_folder_files_operation_cannot_read(tmp_path):
    inp=tmp_path/'in';inp.mkdir();(inp/'a.pdf').write_bytes(b'%PDF');(inp/'b.docx').write_bytes(b'x');(inp/'c.mp4').write_bytes(b'x')
    out=tmp_path/'out';out.mkdir();stats={}
    rows,n=w.scan([str(inp)],out.resolve(),'pdf_images',stats)
    assert [Path(r['sources'][0]).name for r in rows]==['a.pdf'];assert n==1;assert stats['skipped']==1
    rows,_=w.scan([str(inp)],out.resolve(),'rag');assert len(rows)==2

def test_threaded_render_keeps_every_page_in_order(tmp_path):
    pdf=tmp_path/'d.pdf';c=canvas.Canvas(str(pdf))
    for n in range(12):c.drawString(50,700,f'p{n}');c.showPage()
    c.save();work=tmp_path/'w';work.mkdir()
    w.THREADS=4
    try:files=w.convert([pdf],work,{'operation':'pdf_images','dpi':72,'format':'png','range':'','quality':90})
    finally:w.THREADS=1
    assert [f.name for f in files]==[f'd.pdf_page{n:05d}.png' for n in range(1,13)]

def test_images_pdf_accepts_webp_with_alpha(tmp_path):
    a=tmp_path/'a.webp';Image.new('RGBA',(40,30),(0,0,255,100)).save(a)
    b=tmp_path/'b.png';Image.new('RGB',(40,30),'red').save(b)
    work=tmp_path/'w';work.mkdir()
    files=w.convert([a,b],work,{'operation':'images_pdf','paper':'original'})
    from pypdf import PdfReader
    assert len(files)==1 and len(PdfReader(files[0]).pages)==2

def test_wps_office_extensions_are_accepted():
    assert {'.wps','.et','.dps','.docm','.csv'}<=w.OFFICE

def test_clean_naming_page_folder_and_collisions(tmp_path,cfg):
    cfg['naming']='clean'
    src=pdf(tmp_path/'讲义.pdf',n=3);out=tmp_path/'out'
    r=run(src,out,cfg)
    assert sorted(Path(x['path']).relative_to(out).as_posix() for x in r['outputs'])==['讲义/讲义_001.png','讲义/讲义_002.png','讲义/讲义_003.png']
    assert run(src,out,cfg)['state']=='resumed'
    cfg['dpi']=96  # different settings: a new folder, earlier results untouched
    r2=run(src,out,cfg)
    assert {Path(x['path']).parent.name for x in r2['outputs']}=={'讲义 (2)'}
    assert len(list((out/'讲义').iterdir()))==3

def test_clean_naming_single_file_collision(tmp_path):
    out=tmp_path/'out';out.mkdir();(out/'a.jpg').write_bytes(b'unrelated')
    src=tmp_path/'a.png';Image.new('RGB',(20,10),'red').save(src)
    cfg={'operation':'image_convert','format':'jpeg','quality':90,'naming':'clean','lossy':False,'max_edge':0}
    r=run(src,out,cfg)
    assert Path(r['outputs'][0]['path']).name=='a (2).jpg' and (out/'a.jpg').read_bytes()==b'unrelated'

def test_render_reports_page_progress(tmp_path):
    src=pdf(tmp_path/'p.pdf',n=7);work=tmp_path/'w';work.mkdir();prog=tmp_path/'prog.json'
    w.PROGRESS=str(prog);w.THREADS=2
    try:w.convert([src],work,{'operation':'pdf_images','dpi':72,'format':'png','range':'2-6','quality':90})
    finally:w.PROGRESS=None;w.THREADS=1
    assert json.loads(prog.read_text())=={'done':5,'total':5}

def test_plan_units_groups_small_images_only(tmp_path):
    imgs=[]
    for i in range(10):
        p=tmp_path/f'{i}.png';Image.new('RGB',(4,4)).save(p);imgs.append({'sources':[str(p)],'config':{'operation':'image_convert'}})
    units=w.plan_units(imgs,2)
    assert sum(len(u) for u in units)==10 and max(len(u) for u in units)==2
    assert [i for u in units for i,_ in u]==list(range(10))
    pdfs=[{'sources':[str(pdf(tmp_path/'a.pdf'))],'config':{'operation':'pdf_images'}}]*3
    assert all(len(u)==1 for u in w.plan_units(pdfs,2))

def test_chunk_isolates_a_bad_file(tmp_path):
    inp=tmp_path/'in';inp.mkdir()
    for i in range(6):Image.new('RGB',(8,8),'blue').save(inp/f'{i}.png')
    (inp/'3.png').write_bytes(b'not an image')
    req=tmp_path/'r.json';req.write_text(json.dumps({'inputs':[str(inp)],'output':str(tmp_path/'out'),'config':{'operation':'image_convert','format':'jpeg','quality':80,'workers':1,'naming':'clean'}}))
    out=subprocess.run([sys.executable,'-I','-B',str(ROOT/'engine/workbench.py'),'--request',str(req)],capture_output=True,text=True,encoding='utf-8',stdin=subprocess.DEVNULL).stdout
    fin=json.loads(out.splitlines()[-1])
    assert fin['type']=='finished' and fin['success']==5 and fin['failed']==1
