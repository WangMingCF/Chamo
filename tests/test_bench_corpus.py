"""The benchmark corpus has to be reproducible, or the benchmark is not one.

A comparison is only meaningful if the input did not change between two runs, and
across three CI platforms it is only meaningful if the same content produces the
same bytes. Both claims are tested here:

  * two builds of the same suite are byte-identical (and the failing zip member is
    named when they are not);
  * the archive metadata that CPython would otherwise take from the host - entry
    time, "version made by", compression - is pinned, so a Windows runner and a
    Linux runner cannot disagree about the same document;
  * the generated Office files are really openable documents, not just valid bytes.

The suite digest is published to the job summary and to results/ so the three
platforms can be compared after the run; a digest that differs between runners is
the signal that one of the claims above stopped holding.
"""
from __future__ import annotations

import hashlib
import importlib.util
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def _load(name: str, relative: str):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


corpus = _load("chamo_bench_corpus", "scripts/bench_corpus.py")

ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)


def _member_differences(left: Path, right: Path) -> list:
    """Which parts inside two OOXML packages differ (empty for a non-zip file).

    A byte difference in a .xlsx or .pptx says "something changed" but not what,
    and the writer that changed it is what the next fix has to target. Naming the
    members turns a red test into a diagnosis.
    """
    import zipfile

    try:
        with zipfile.ZipFile(left) as first, zipfile.ZipFile(right) as second:
            names = sorted(set(first.namelist()) | set(second.namelist()))
            changed = []
            for name in names:
                if name not in first.namelist() or name not in second.namelist():
                    changed.append(name + " (present on one side only)")
                elif first.read(name) != second.read(name):
                    changed.append(name)
            return changed
    except zipfile.BadZipFile:
        return []


def _digest(manifest: dict) -> str:
    joined = "\n".join("%s %s" % (entry["path"], entry["sha256"])
                       for entry in sorted(manifest["files"], key=lambda e: e["path"]))
    return hashlib.sha256(joined.encode("utf-8")).hexdigest()


def publish_digest(manifest: dict) -> str:
    """Publish the suite digest so the three runners' outputs can be compared.

    Reproducibility across machines cannot be checked from inside one job; it is
    checked by comparing this value between the platform summaries. When the job
    runs outside Actions the value is only printed.
    """
    import os

    digest = _digest(manifest)
    payload = {"suite": manifest["suite"], "digest": digest,
               "files": {entry["path"]: entry["sha256"] for entry in manifest["files"]}}
    print(("corpus %s digest: %s" % (manifest["suite"], digest)))
    if os.environ.get("GITHUB_ACTIONS") == "true":
        artifacts = Path("results")
        artifacts.mkdir(parents=True, exist_ok=True)
        (artifacts / ("bench-corpus-%s.json" % manifest["suite"])).write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    summary_path = os.environ.get("GITHUB_STEP_SUMMARY")
    if summary_path:
        try:
            with open(summary_path, "a", encoding="utf-8") as handle:
                handle.write("### Benchmark corpus\n\n"
                             "| suite | files | digest |\n| --- | --- | --- |\n"
                             "| %s | %d | %s |\n\n"
                             % (manifest["suite"], len(manifest["files"]), digest))
        except OSError as error:
            print("could not append to GITHUB_STEP_SUMMARY (%s)" % error)
    return digest


def test_the_same_suite_builds_the_same_bytes(tmp_path):
    """A benchmark comparison is only meaningful if the input did not change."""
    first = corpus.build(tmp_path / "a", "smoke")
    second = corpus.build(tmp_path / "b", "smoke")
    differing = [entry["path"] for entry, other in zip(first["files"], second["files"])
                 if entry["sha256"] != other["sha256"]]
    detail = {name: _member_differences(tmp_path / "a" / name, tmp_path / "b" / name)
              for name in differing}
    assert differing == [], "the smoke corpus is not reproducible: %s" % detail
    assert corpus.check(tmp_path / "a") == []
    publish_digest(first)


def test_archive_metadata_is_pinned_not_taken_from_the_host(tmp_path):
    """The three fields CPython fills in from the host must be constants.

    Without this the "same bytes on every machine" claim is only true on one
    platform: create_system is 0 on Windows and 3 elsewhere, the entry time is the
    current time, and compressed entries depend on the zlib build.
    """
    import zipfile

    corpus.build(tmp_path / "c", "smoke")
    for name in ("memo.docx", "deck.pptx", "sheet.xlsx"):
        with zipfile.ZipFile(tmp_path / "c" / name) as archive:
            for info in archive.infolist():
                assert info.date_time == ZIP_TIMESTAMP, (name, info.filename, info.date_time)
                assert info.create_system == 3, (name, info.filename, info.create_system)
                assert info.compress_type == zipfile.ZIP_STORED, (name, info.filename)


def test_the_office_files_are_real_documents(tmp_path):
    """Valid bytes are not the same as a document an application can open."""
    import xml.etree.ElementTree as ET
    import zipfile

    from openpyxl import load_workbook
    from pptx import Presentation

    corpus.build(tmp_path / "d", "smoke")
    workbook = load_workbook(tmp_path / "d" / "sheet.xlsx")
    assert workbook.sheetnames == ["Numbers", "Summary"]
    assert workbook["Numbers"]["B2"].value is not None

    presentation = Presentation(tmp_path / "d" / "deck.pptx")
    assert len(presentation.slides) == corpus.SUITES["smoke"]["slides"]

    with zipfile.ZipFile(tmp_path / "d" / "memo.docx") as archive:
        document = ET.fromstring(archive.read("word/document.xml"))
    text = "".join(node.text or "" for node in document.iter())
    assert "Paragraph 1" in text


def test_a_changed_or_extra_file_is_reported(tmp_path):
    corpus.build(tmp_path / "c", "smoke")
    assert corpus.check(tmp_path / "c") == []
    (tmp_path / "c" / "text.pdf").write_bytes(b"not the same pdf")
    (tmp_path / "c" / "stray.bin").write_bytes(b"x")
    issues = corpus.check(tmp_path / "c")
    assert any("text.pdf" in issue and "changed" in issue for issue in issues), issues
    assert any("stray.bin" in issue for issue in issues), issues


def test_a_damaged_manifest_is_reported_not_crashed(tmp_path):
    corpus.build(tmp_path / "e", "smoke")
    (tmp_path / "e" / "manifest.json").write_text("{not json", encoding="utf-8")
    issues = corpus.check(tmp_path / "e")
    assert any("cannot be read" in issue for issue in issues), issues


def test_refuses_to_write_over_a_previous_corpus(tmp_path):
    corpus.build(tmp_path / "d", "smoke")
    try:
        corpus.build(tmp_path / "d", "smoke")
    except SystemExit as error:
        assert "non-empty" in str(error)
    else:
        raise AssertionError("building into a used directory must be refused")


def test_an_interrupted_build_does_not_lock_the_directory(tmp_path):
    """A leftover .normalizing file is not corpus content and must not block a retry."""
    (tmp_path / "f").mkdir()
    (tmp_path / "f" / "memo.docx.normalizing").write_bytes(b"half-written")
    corpus.build(tmp_path / "f", "smoke")
    assert corpus.check(tmp_path / "f") == []
