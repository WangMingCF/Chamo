import csv
import datetime
import itertools
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from .output import Sink


def serialise(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    if isinstance(value, datetime.timedelta):
        return {"type": "timedelta", "value": value.total_seconds()}
    return value


def deserialise(value):
    if not isinstance(value, dict):
        return value
    if value["type"] == "timedelta":
        return datetime.timedelta(seconds=value["value"])
    return getattr(datetime, value["type"]).fromisoformat(value["value"])


def pack(source: Path, work: Path, sink: Sink, options: dict):
    header_count = options["header_rows"]
    if header_count >= sink.limits.table_rows:
        raise ValueError("header_rows 必须小于 table_rows，为正文留出至少一行")
    sink.notice("表格按工作表拆成 XLSX 数据快照：不保留样式、图表、图片、合并布局或宏；首行默认作为重复表头，可用 --header-rows 0 关闭")
    sink.notice("公式以文本保留，不执行/重算公式；跨分片引用不保证可计算")
    book = None
    csv_handle = None
    try:
        if source.suffix.lower() == ".csv":
            csv_handle = source.open(encoding=options["encoding"], newline="")
            # CSV retains leading zeros and protects cells beginning with '=' as literal text.
            csv.field_size_limit(options["csv_field_bytes"])
            sheets = [("Sheet1", ([{"v": serialise(v), "format": "General", "kind": "s"} for v in row] for row in csv.reader(csv_handle)))]
        else:
            book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
            def rows(sheet):
                sheet.reset_dimensions()  # Do not trust stale OOXML used-range metadata.
                for row in sheet.iter_rows():
                    values = []
                    for cell in row:
                        value = cell.value
                        if cell.data_type == "f" and not isinstance(value, str):
                            raise ValueError("不支持数组/数据表公式；请先在 Excel 中转为值")
                        values.append({"v": serialise(value), "format": cell.number_format,
                                       "kind": "s" if isinstance(value, str) else cell.data_type})
                    yield values
            sheets = ((sheet.title, rows(sheet)) for sheet in book)
        for sheet_name, iterator in sheets:
            iterator = iter(iterator)
            headers = list(itertools.islice(iterator, header_count))
            first_row = len(headers) + 1
            if not headers and not header_count:
                first_row = 1
            capacity = sink.limits.table_rows - len(headers)
            emitted_sheet = False
            while True:
                spool = work / "rows.jsonl"
                count = 0
                with spool.open("w", encoding="utf-8") as out:
                    for row in itertools.islice(iterator, capacity):
                        out.write(json.dumps(row, ensure_ascii=False) + "\n")
                        count += 1
                if not count and emitted_sheet:
                    break
                if not count and not headers:
                    sink.notice(f"空工作表 {sheet_name!r} 没有数据，未生成空附件")
                    break

                def emit_range(start, end):
                    output = Workbook(write_only=True)
                    sheet = output.create_sheet(sheet_name)
                    def append(row):
                        cells = []
                        for item in row:
                            value = deserialise(item["v"])
                            if isinstance(value, str) and len(value) > 32767:
                                # openpyxl otherwise truncates silently at Excel's cell limit.
                                raise ValueError("单个字段超过 XLSX 的 32767 字符上限，无法无损写入；请拆分字段后再处理")
                            cell = WriteOnlyCell(sheet, value=value)
                            # Text stays text: a formula or CSV string must never execute.
                            if item["kind"] == "s":
                                cell.data_type = "s"
                            cell.number_format = item["format"]
                            cells.append(cell)
                        sheet.append(cells)
                    for row in headers:
                        append(row)
                    with spool.open(encoding="utf-8") as handle:
                        for line in itertools.islice(handle, start, end):
                            append(json.loads(line))
                    candidate = work / "table.xlsx"
                    output.save(candidate)
                    output.close()
                    if candidate.stat().st_size <= sink.limits.table_bytes:
                        sink.emit(candidate, "table", {"sheet": sheet_name,
                                  "row_start": first_row + start if end > start else None,
                                  "row_end": first_row + end - 1 if end > start else None,
                                  "header_rows": len(headers)})
                    elif end - start > 1:
                        mid = (start + end) // 2
                        emit_range(start, mid)
                        emit_range(mid, end)
                    else:
                        raise ValueError(f"工作表 {sheet_name!r} 单行加表头仍超过 table_bytes，请减少表头或提高限制")
                emit_range(0, count)
                emitted_sheet = True
                first_row += count
                if count < capacity:
                    break
    finally:
        if book:
            book.close()
        if csv_handle:
            csv_handle.close()
