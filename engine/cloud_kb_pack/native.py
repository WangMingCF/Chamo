"""Conservative native-file handling for RAG: never flatten content by default."""
import zipfile
from lxml import etree
import olefile


def xml(stream):
    return etree.parse(stream, etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False))


def inspect_office(path):
    suffix = path.suffix.lower()
    if olefile.isOleFile(str(path)):
        with olefile.OleFileIO(str(path)) as container:
            if container.exists("EncryptedPackage") or container.exists("EncryptionInfo"):
                raise ValueError("加密 Office 文件，未提供密码，保留原件等待处理")
            if suffix in {".doc", ".docx"}:
                if not container.exists("WordDocument"):
                    raise ValueError("Word OLE 容器缺少 WordDocument 流")
                with container.openstream("WordDocument") as stream:
                    header = stream.read(32)
                    if len(header) >= 12 and int.from_bytes(header[10:12], "little") & 0x8100:
                        raise ValueError("加密 Word 文件，保留原件等待处理")
            if suffix == ".ppt" and not container.exists("PowerPoint Document"):
                raise ValueError("PPT OLE 容器缺少演示文稿流")
            if suffix == ".xls" and not (container.exists("Workbook") or container.exists("Book")):
                raise ValueError("XLS OLE 容器缺少工作簿流")
            metadata = container.get_metadata()
            return {"pages_metadata": metadata.slides if suffix == ".ppt" else metadata.num_pages,
                    "page_count_method": "source_metadata_unverified", "native": "ole"}
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            bad = archive.testzip()
            if bad:
                raise ValueError(f"Office ZIP 校验失败：{bad}")
            if suffix in {".doc", ".docx"}:
                with archive.open("word/document.xml") as handle:
                    for _, node in etree.iterparse(handle, events=("end",), resolve_entities=False, no_network=True):
                        node.clear()
                        while node.getprevious() is not None:
                            del node.getparent()[0]
            pages = None
            if suffix in {".ppt", ".pptx"}:
                with archive.open("ppt/presentation.xml") as handle:
                    pages = len(xml(handle).findall(".//{http://schemas.openxmlformats.org/presentationml/2006/main}sldId"))
            elif "docProps/app.xml" in archive.namelist():
                with archive.open("docProps/app.xml") as handle:
                    node = xml(handle).find(".//{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Pages")
                    if node is not None and node.text and node.text.isdecimal():
                        pages = int(node.text)
            return {"pages_metadata": pages, "page_count_method": "slide_list" if suffix in {".ppt", ".pptx"} else "source_metadata_unverified", "native": "ooxml"}
    with path.open("rb") as handle:
        header = handle.read(65536).lstrip()
    if header.startswith(b"{\\rtf") or (suffix == ".doc" and header.lower().startswith((b"<html", b"<!doctype html"))):
        return {"pages_metadata": None, "page_count_method": "unknown", "native": "rtf_or_html"}
    if suffix == ".doc":
        # Word can save an MHTML web archive with a .doc extension. Its related MIME
        # parts include the original HTML and images; keep them as one native file.
        from email.parser import BytesHeaderParser
        from email import policy
        mime = BytesHeaderParser(policy=policy.default).parsebytes(header)
        if mime.get_content_type() == "multipart/related" and mime.get_boundary():
            from email.parser import BytesFeedParser
            parser = BytesFeedParser(policy=policy.default)
            with path.open("rb") as stream:
                while block := stream.read(65536):
                    parser.feed(block)
            message = parser.close()
            if not message.is_multipart() or any(part.defects for part in message.walk()):
                raise ValueError("Word MHTML MIME 结构损坏，保留原件")
            if not any(part.get_content_type() == "text/html" for part in message.walk()):
                raise ValueError("Word MHTML 缺少 HTML 正文")
            return {"pages_metadata": None, "page_count_method": "unknown", "native": "mhtml_word"}
    raise ValueError("Office 文件损坏或格式不一致，保留原件等待处理")


def inspect_table(path):
    if path.suffix.lower() == ".xls":
        facts = inspect_office(path)
        import xlrd
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            rows = sum(workbook.sheet_by_index(index).nrows for index in range(workbook.nsheets))
        finally:
            workbook.release_resources()
        return {**facts, "rows": rows}
    if path.suffix.lower() == ".csv":
        return {}
    from openpyxl import load_workbook
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise ValueError(f"XLSX ZIP 校验失败：{bad}")
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        total = 0
        for sheet in book:
            sheet.reset_dimensions()
            total += sum(1 for _ in sheet.iter_rows())
        return {"rows": total, "native": "xlsx"}
    finally:
        book.close()


def pack(source, work, sink, options, category):
    if category == "document":
        facts = inspect_office(source)
        pages = facts["pages_metadata"]
        if pages is not None and pages > sink.limits.document_pages:
            raise ValueError("Office 记录的页数超限；不能确认无损拆分，保留原件")
        if source.stat().st_size > sink.limits.document_bytes:
            if source.suffix.lower() == ".docx":
                from .docx import split
                split(source, work, sink, options)
                return
            raise ValueError("原生 Office 超限，不能确认无损拆分，保留原件；不转换为 PDF")
        if facts["page_count_method"] != "slide_list":
            sink.notice("Word 原样保留；页数来自源元数据，未重新排版核验，未知页数不会被宣称已验证")
        sink.emit(source, "document", {"operation": "copy_native", **facts}, verified=facts)
    elif category == "table":
        if source.suffix.lower() == ".csv":
            import csv
            with source.open(encoding=options["encoding"], newline="") as handle:
                rows = sum(1 for _ in csv.reader(handle))
            facts = {"rows": rows, "native": "csv"}
        else:
            facts = inspect_table(source)
        if facts["rows"] > sink.limits.table_rows or source.stat().st_size > sink.limits.table_bytes:
            raise ValueError("表格超限，不能确认公式、合并、图表和跨表关系无损拆分，保留原件")
        sink.emit(source, "table", {"operation": "copy_native", **facts}, verified=facts)
    elif category == "image":
        from PIL import Image
        if source.suffix.lower() not in {".png", ".jpg", ".jpeg", ".bmp", ".gif"}:
            raise ValueError("图片格式不在目标支持列表；不能默认有损转换，保留原件")
        with Image.open(source) as image:
            w, h = image.size
            image.verify()
        if min(w, h) < sink.limits.image_min_edge or max(w, h) > sink.limits.image_max_edge or max(w, h)/min(w,h) >= sink.limits.image_max_ratio:
            raise ValueError("图片超限；禁止降分辨率、裁剪或有损压缩，保留原件")
        if source.stat().st_size > sink.limits.image_bytes:
            from . import lossless
            if source.suffix.lower() == ".png":
                target, proof = lossless.png(source, work)
            elif source.suffix.lower() in {".jpg", ".jpeg"}:
                target, proof = lossless.jpeg(source, work)
            else:
                raise ValueError("图片超限；没有经过验证的无损压缩方案，保留原件")
            if target.stat().st_size > sink.limits.image_bytes:
                raise ValueError("图片无损压缩后仍超限；禁止有损压缩，保留原件")
            sink.emit(target, "image", {"operation": "lossless_compress", "proof": proof})
            return
        sink.emit(source, "image", {"operation": "copy_native", "frames_preserved": True}, verified={"width": w, "height": h})
